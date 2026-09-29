"""Independent verification of Nifty 500 Value 50 (v50) price/TRI data.
Compares stored /app/build/data_snapshots/latest.json against:
 (a) official niftyindices.com API
 (b) Excel workbook /app/NSE_Index_Data_Price_TRI.xlsx
"""
import json
import os
from datetime import datetime

import pytest
import requests
from openpyxl import load_workbook

SNAP = "/app/build/data_snapshots/latest.json"
XLSX = "/app/NSE_Index_Data_Price_TRI.xlsx"
TOL = 0.05

SAMPLE_DATES = [
    "2005-04-01",  # base
    "2010-01-04",
    "2015-06-01",
    "2020-03-23",
    "2023-12-29",
    "2026-09-28",  # latest
]

EXPECTED_LATEST = {"price": 15341.50, "tri": 26160.57}
EXPECTED_BASE = {"price": 1000.00, "tri": 1000.00}


@pytest.fixture(scope="module")
def snap():
    with open(SNAP) as f:
        return json.load(f)


@pytest.fixture(scope="module")
def stored_by_date(snap):
    dates = snap["dates"]
    price = snap["price"]["v50"]
    tri = snap["tri"]["v50"]
    return {d: (price[i], tri[i]) for i, d in enumerate(dates)}


def _fmt(d):
    return datetime.strptime(d, "%Y-%m-%d").strftime("%d-%b-%Y")


@pytest.fixture(scope="module")
def official_data():
    """Fetch official Price + TRI history from niftyindices.com."""
    sess = requests.Session()
    ua = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
          "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36")
    sess.headers.update({
        "User-Agent": ua,
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Accept-Language": "en-US,en;q=0.9",
        "Referer": "https://www.niftyindices.com/reports/historical-data",
        "Origin": "https://www.niftyindices.com",
        "X-Requested-With": "XMLHttpRequest",
    })
    try:
        sess.get("https://www.niftyindices.com/reports/historical-data", timeout=20)
    except Exception as e:
        pytest.skip(f"Cannot reach niftyindices.com: {e}")

    body_tmpl = ("{{\"cinfo\":\"{{'name':'Nifty500 Value 50',"
                 "'startDate':'01-Apr-2005','endDate':'30-Sep-2026',"
                 "'indexName':'Nifty500 Value 50'}}\"}}")
    body = body_tmpl.format()
    headers = {"Content-Type": "application/json; charset=utf-8"}

    def _post(url):
        r = sess.post(url, data=body.encode("utf-8"), headers=headers, timeout=60)
        r.raise_for_status()
        txt = r.content.decode("utf-8-sig")
        return json.loads(txt)

    try:
        price_rows = _post("https://www.niftyindices.com/BackPage/getHistoricaldatatabletoString")
        tri_rows = _post("https://www.niftyindices.com/BackPage/getTotalReturnIndexString")
    except Exception as e:
        pytest.skip(f"niftyindices API failed: {e}")

    price_map = {}
    for row in price_rows:
        d = datetime.strptime(row["HistoricalDate"], "%d %b %Y").strftime("%Y-%m-%d")
        price_map[d] = float(row["CLOSE"])
    tri_map = {}
    for row in tri_rows:
        d = datetime.strptime(row["Date"], "%d %b %Y").strftime("%Y-%m-%d")
        tri_map[d] = float(row["TotalReturnsIndex"])
    return {"price": price_map, "tri": tri_map}


@pytest.fixture(scope="module")
def excel_data():
    wb = load_workbook(XLSX, data_only=True, read_only=True)
    out = {"price": {}, "tri": {}}
    mapping = {"Price (Unadjusted)": "price", "TRI (Total Return)": "tri"}
    for sheet_name, key in mapping.items():
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        try:
            col_idx = header.index("Nifty 500 Value 50")
        except ValueError:
            # try alternate name
            col_idx = header.index("Nifty500 Value 50")
        date_col = 0
        for row in rows:
            dv = row[date_col]
            v = row[col_idx]
            if dv is None or v is None:
                continue
            if isinstance(dv, datetime):
                d = dv.strftime("%Y-%m-%d")
            else:
                s = str(dv).strip()
                d = None
                for fmt in ("%Y-%m-%d", "%d-%b-%Y", "%d-%m-%Y", "%d/%m/%Y"):
                    try:
                        d = datetime.strptime(s, fmt).strftime("%Y-%m-%d")
                        break
                    except Exception:
                        pass
                if d is None:
                    continue
            try:
                out[key][d] = float(v)
            except Exception:
                pass
    return out


# ---------- Basic sanity ----------

def test_stored_length(snap):
    assert len(snap["dates"]) == len(snap["price"]["v50"]) == len(snap["tri"]["v50"])


def test_stored_base(stored_by_date):
    p, t = stored_by_date["2005-04-01"]
    assert abs(p - EXPECTED_BASE["price"]) < TOL
    assert abs(t - EXPECTED_BASE["tri"]) < TOL


def test_stored_latest(stored_by_date):
    p, t = stored_by_date["2026-09-28"]
    assert abs(p - EXPECTED_LATEST["price"]) < TOL, f"price {p}"
    assert abs(t - EXPECTED_LATEST["tri"]) < TOL, f"tri {t}"


# ---------- Compare stored vs official API ----------

@pytest.mark.parametrize("date", SAMPLE_DATES)
def test_stored_matches_official(date, stored_by_date, official_data):
    if date not in stored_by_date:
        pytest.skip(f"date {date} not in stored series")
    sp, st = stored_by_date[date]
    op = official_data["price"].get(date)
    ot = official_data["tri"].get(date)
    assert op is not None, f"official price missing for {date}"
    assert ot is not None, f"official tri missing for {date}"
    assert abs(sp - op) <= TOL, f"PRICE mismatch {date}: stored={sp} official={op}"
    assert abs(st - ot) <= TOL, f"TRI mismatch {date}: stored={st} official={ot}"


# ---------- Compare stored vs Excel ----------

@pytest.mark.parametrize("date", SAMPLE_DATES)
def test_stored_matches_excel(date, stored_by_date, excel_data):
    if date not in stored_by_date:
        pytest.skip(f"{date} not in stored")
    sp, st = stored_by_date[date]
    ep = excel_data["price"].get(date)
    et = excel_data["tri"].get(date)
    if ep is None or et is None:
        pytest.skip(f"{date} not in Excel workbook")
    assert abs(sp - ep) <= TOL, f"PRICE stored={sp} excel={ep}"
    assert abs(st - et) <= TOL, f"TRI stored={st} excel={et}"
